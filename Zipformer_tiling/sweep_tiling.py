import json
import os
import math
from datetime import datetime
import numpy as np

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("❌ LỖI: Cậu chưa cài thư viện openpyxl. Hãy chạy lệnh: pip install openpyxl")
    exit()

from zipformer_modular_accelerator import (
    HardwareConfig, FeedforwardHardware, FeedforwardChannelFirst, FeedforwardBlockTiling,
    NonlinearAttentionHardware, NonlinearAttentionChannelFirst, NonlinearAttentionBlockTiling, 
    SelfAttentionPipelineHardware
)

# CẤU HÌNH KIỂU DỮ LIỆU: 2 cho FP16/INT16, 1 cho INT8 (Lượng tử hóa)
BYTES_PER_WORD = 2 

def get_true_strategy_name(t_tile, d_tile, T, D):
    if t_tile >= T and d_tile >= D: return "Full-Matrix"
    if t_tile >= T: return "Channel-First"
    if d_tile >= D: return "Time-First"
    return "Block Tiling"

# =====================================================================
# 🎯 1. TÍNH DUNG LƯỢNG SRAM CHUẨN VẬT LÝ
# =====================================================================
def calculate_exact_hardware_sram_kb(module_type, strat, t_tile, d_tile, T, D_in, D_out, ping_pong=True):
    P = 2 if ping_pong else 1 # Double Buffering giấu Latency
    
    if module_type == "FFW":
        buf_res = T * D_in # Residual Buffer ngâm suốt cả khối
        buf_x = (T * D_in) if strat in ["Full-Matrix", "Channel-First"] else (P * t_tile * D_in)
        
        if strat in ["Full-Matrix", "Time-First"]:
            buf_w = (D_in * D_out) + (D_out * D_in)
        else: # Channel-First hoặc Block Tiling
            buf_w = P * ((D_in * d_tile) + (d_tile * D_in))
            
        buf_out = (T * d_tile) if strat == "Channel-First" else (t_tile * d_tile)
        total_words = buf_res + buf_x + buf_w + buf_out
        
    elif module_type == "NONLIN":
        buf_aux = (T * D_in) + (T * T) # Residual + Attn Map (T x T)
        buf_x = (T * D_in) if strat in ["Full-Matrix", "Channel-First"] else (P * t_tile * D_in)
        buf_w = (D_in * D_out) if strat in ["Full-Matrix", "Time-First"] else (P * D_in * d_tile)
        buf_fusion_out = (T * d_tile) if strat == "Channel-First" else (t_tile * d_tile)
        
        total_words = buf_aux + buf_x + buf_w + buf_fusion_out

    else: # SELF_ATTENTION
        vhd = d_tile // 4
        total_words = (T * D_in) + (D_in * vhd) + (T * vhd) * 2 + (T * vhd) + (T * d_tile) + (T * D_in) + (T * T)

    return (total_words * BYTES_PER_WORD) / 1024.0

# =====================================================================
# 🎯 2. TÍNH CHÍNH XÁC SỐ LẦN LOADS / STORES TỪ DRAM
# =====================================================================
def calculate_exact_dram_traffic(module_type, strat, t_tile, d_tile, T, D_in, D_out):
    N_T = math.ceil(T / t_tile)
    N_C = math.ceil(D_out / d_tile)

    if module_type == "FFW":
        if strat in ["Full-Matrix", "Channel-First"]:
            g_loads = (T * D_in) + (D_in * D_out) + (D_out * D_in)
        elif strat == "Time-First":
            g_loads = (T * D_in) + N_T * ((D_in * D_out) + (D_out * D_in))
        else: # Block Tiling
            g_loads = N_C * (T * D_in) + N_T * ((D_in * D_out) + (D_out * D_in))
        
        g_stores = T * D_in

    elif module_type == "NONLIN":
        if strat in ["Full-Matrix", "Channel-First"]:
            g_loads = (T * D_in) + (D_in * D_out)
        elif strat == "Time-First":
            g_loads = (T * D_in) + N_T * (D_in * D_out)
        else: # Block Tiling
            g_loads = N_C * (T * D_in) + N_T * (D_in * D_out)
            
        g_stores = T * D_in

    else: # SELF_ATTENTION PIPELINE
        vhd = D_out // 4
        g_loads = (T * D_in) + (D_in * vhd) + (T * T)
        g_stores = T * D_in

    return g_loads, g_stores

def calculate_latency_ms(macs, total_traffic_words, num_pe=256, freq_mhz=200, bw_gbps=10):
    cycles_compute = math.ceil(macs / num_pe)
    time_compute_sec = cycles_compute / (freq_mhz * 1e6)
    time_memory_sec = (total_traffic_words * BYTES_PER_WORD) / (bw_gbps * 1e9)
    return (time_compute_sec + time_memory_sec) * 1000

def run_multi_stack_dse():
    print("\n" + "★"*165)
    print(f"🚀 DSE CHUẨN XÁC: ĐÃ MỞ RỘNG TILE MẢNG VÀ HOÀN THIỆN XỬ LÝ STACK 3 (Data Precision: {BYTES_PER_WORD*8}-bit) 🚀")
    print("★"*165)
    
    zipformer_stacks = [
        (0, 116, 192,  384,  432,  144,  48),
        (1, 58,  384,  768,  864,  288,  96),
        (2, 29,  768,  1536, 1728, 576,  192),
        (3, 14,  1536, 3072, 3456, 1152, 384),
        (4, 29,  768,  1536, 1728, 576,  192),
        (5, 58,  384,  768,  864,  288,  96)
    ]

    ffw_classes = {"Full-Matrix": FeedforwardHardware, "Time-First": FeedforwardHardware, "Channel-First": FeedforwardChannelFirst, "Block Tiling": FeedforwardBlockTiling}
    nonlin_classes = {"Full-Matrix": NonlinearAttentionHardware, "Time-First": NonlinearAttentionHardware, "Channel-First": NonlinearAttentionChannelFirst, "Block Tiling": NonlinearAttentionBlockTiling}
    
    total_system_latency_ms = 0.0
    golden_configs_export, excel_rows_export = [], []
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    HW_PE, HW_FREQ, HW_BW = 256, 200, 10

    header = f"{'Stk':<3} | {'Shape(T,D)':<12} | {'FFW Strat':<14} | {'FFW Tile':<11} | {'Nonlin Strat':<14} | {'Nonlin Tile':<12} | {'Self Tile':<10} | {'Exact SRAM':<11} | {'Sys M/Byte':<10} | {'Latency':<10}"
    print(header)
    print("-" * 165)

    for stack in zipformer_stacks:
        s_id, T, D_in, D_ffw, D_attn, D_cn, D_cs = stack
        
        t_factors = [1, 2, 4]
        # [SỬA QUAN TRỌNG]: Mở rộng d_factors lên 256 để chứa các Tile nhỏ cho Stack 3!
        d_factors = [1, 2, 4, 6, 8, 12, 16, 24, 32, 48, 64, 96, 128, 192, 256] 

        f_options, n_options = [], []
        for tf in t_factors:
            t_cut = max(1, T // tf)
            for df in d_factors:
                f_h = max(16, int(math.ceil(D_ffw / df)))
                
                # Ép n_c chia hết cho 3 cho khối Nonlinear Attention
                n_c_raw = max(18, int(math.ceil(D_attn / df)))
                n_c = math.ceil(n_c_raw / 3.0) * 3
                
                f_options.append((t_cut, f_h))
                n_options.append((t_cut, n_c))

        f_options = list(set(f_options))
        n_options = list(set(n_options))

        best_intensity = -1
        best_candidate = None

        # =====================================================================
        # ⚡ BƯỚC 1: QUÉT DSE LỌC CẤU HÌNH VÀNG
        # =====================================================================
        for f_t, f_h in f_options:
            for n_t, n_c in n_options:
                f_strat = get_true_strategy_name(f_t, f_h, T, D_ffw)
                n_strat = get_true_strategy_name(n_t, n_c, T, D_attn)
                s_t = max(f_t, n_t)
                
                sram_f = calculate_exact_hardware_sram_kb("FFW", f_strat, f_t, f_h, T, D_in, D_ffw, ping_pong=True)
                sram_n = calculate_exact_hardware_sram_kb("NONLIN", n_strat, n_t, n_c, T, D_in, D_attn, ping_pong=True)
                sram_s = calculate_exact_hardware_sram_kb("SELF", "Pipeline", s_t, D_cs, T, D_in, D_cs, ping_pong=False)
                peak_sram = max(sram_f, sram_n, sram_s)

                if peak_sram > 512.0:
                    continue

                macs_f = T * D_in * D_ffw * 2
                macs_n = T * D_in * D_attn * 2
                macs_s = T * D_in * D_cs * 4
                tot_macs = macs_f + macs_n + macs_s

                loads_f, stores_f = calculate_exact_dram_traffic("FFW", f_strat, f_t, f_h, T, D_in, D_ffw)
                loads_n, stores_n = calculate_exact_dram_traffic("NONLIN", n_strat, n_t, n_c, T, D_in, D_attn)
                loads_s, stores_s = calculate_exact_dram_traffic("SELF", "Pipeline", s_t, D_cs, T, D_in, D_cs)

                tot_loads = loads_f + loads_n + loads_s
                tot_stores = stores_f + stores_n + stores_s
                total_traffic_words = tot_loads + tot_stores

                intensity_mac_byte = tot_macs / (total_traffic_words * BYTES_PER_WORD) if total_traffic_words > 0 else 0

                if intensity_mac_byte > best_intensity:
                    best_intensity = intensity_mac_byte
                    best_candidate = (f_strat, f_t, f_h, n_strat, n_t, n_c, "Pipeline", s_t, peak_sram, tot_macs, tot_loads, tot_stores, intensity_mac_byte)

        # =====================================================================
        # 🧪 BƯỚC 2: CHẠY PROFILER TRÍCH XUẤT XÁC MINH (VERIFICATION)
        # =====================================================================
        if best_candidate:
            b_f_strat, b_f_t, b_f_h, b_n_strat, b_n_t, b_n_c, b_s_strat, b_s_t, exact_sram_kb, tot_macs, analytical_loads, analytical_stores, intensity_mac_byte = best_candidate
            
            X = np.random.randn(T, D_in).astype(np.float64) / 10.0
            res = np.zeros_like(X)
            attn_w_3d = np.random.rand(4, T, T).astype(np.float64)
            attn_w_2d = attn_w_3d[0]

            cfg_f = HardwareConfig(T, D_in, D_ffw, D_attn, D_cn, b_f_strat, b_f_t, b_f_h, 0, 0, 2048, 512)
            cfg_n = HardwareConfig(T, D_in, D_ffw, D_attn, D_cn, b_n_strat, 0, 0, b_n_t, b_n_c, 2048, 512)
            cfg_s = HardwareConfig(T, D_in, D_ffw, D_attn, D_cs, b_s_strat, 0, 0, b_s_t, 0, 2048, 512)

            hw_f = ffw_classes[b_f_strat](cfg_f)
            hw_n = nonlin_classes[b_n_strat](cfg_n)
            hw_s = SelfAttentionPipelineHardware(cfg_s)

            out_f, prof_f = hw_f.execute(X)
            out_n, prof_n = hw_n.execute(out_f, res, attn_w_2d) 
            out_s, prof_s = hw_s.execute(out_n, res, attn_w_3d)

            prof_macs = prof_f.metrics['macs'] + prof_n.metrics['macs'] + prof_s.metrics['macs']
            prof_loads = prof_f.metrics['g_loads'] + prof_n.metrics['g_loads'] + prof_s.metrics['g_loads']
            prof_stores = prof_f.metrics.get('g_stores', 0) + prof_n.metrics.get('g_stores', 0) + prof_s.metrics.get('g_stores', 0)
            
            total_traffic_words = prof_loads + prof_stores
            latency_ms = calculate_latency_ms(prof_macs, total_traffic_words, HW_PE, HW_FREQ, HW_BW)

            shape_str = f"({T}, {D_in})"
            tile_f_str = f"({b_f_t},{b_f_h})"
            tile_n_str = f"({b_n_t},{b_n_c})"
            
            print(f"Stk {s_id:<1} | {shape_str:<12} | {b_f_strat:<14} | {tile_f_str:<11} | {b_n_strat:<14} | {tile_n_str:<12} | ({b_s_t},-)   | {exact_sram_kb:<9.1f} KB | {intensity_mac_byte:<10.2f} | {latency_ms:<8.3f} ms")
            total_system_latency_ms += latency_ms
            
            golden_configs_export.append({
                "stack_id": s_id, "shape": {"T": T, "D_in": D_in, "D_ffw": D_ffw, "D_attn": D_attn, "D_chunk_n": D_cn, "D_chunk_s": D_cs},
                "ffw": {"strat": b_f_strat, "t": b_f_t, "h": b_f_h}, "nonlin": {"strat": b_n_strat, "t": b_n_t, "c": b_n_c}, "self": {"strat": b_s_strat, "t": b_s_t},
                "hardware_specs": {
                    "exact_sram_kb": round(exact_sram_kb, 2), 
                    "g_loads": prof_loads, 
                    "g_stores": prof_stores, 
                    "intensity_mac_byte": round(intensity_mac_byte, 2)
                }
            })
            
            excel_rows_export.append([
                timestamp, s_id, T, D_in, D_ffw, D_attn, HW_PE, HW_FREQ, HW_BW,
                b_f_strat, f"{b_f_t}x{b_f_h}", b_n_strat, f"{b_n_t}x{b_n_c}", b_s_strat, f"{b_s_t}",
                round(exact_sram_kb, 2), prof_macs, prof_loads, prof_stores,
                prof_f.metrics.get('l_loads', 0) + prof_n.metrics.get('l_loads', 0) + prof_s.metrics.get('l_loads', 0), 
                prof_f.metrics.get('l_stores', 0) + prof_n.metrics.get('l_stores', 0) + prof_s.metrics.get('l_stores', 0),
                round(intensity_mac_byte, 2), round(latency_ms, 4)
            ])
        else:
            print(f"Stk {s_id:<2} | ❌ LỖI: Không tìm thấy cấu hình nào thỏa mãn SRAM <= 512KB!")

    print("-" * 165)
    print(f"⏱️  TỔNG THỜI GIAN SUY LUẬN TỰ ĐỘNG LÀM TRÒN (LATENCY): {total_system_latency_ms:.3f} mili-giây")
    print("★"*165 + "\n")

    script_dir = os.path.dirname(os.path.abspath(__file__))
    
    with open(os.path.join(script_dir, 'golden_configs.json'), 'w') as f:
        json.dump(golden_configs_export, f, indent=4)
        
    excel_path = os.path.join(script_dir, 'dse_results_history.xlsx')
    if os.path.isfile(excel_path):
        wb = load_workbook(excel_path)
        ws = wb.active
    else:
        wb, ws = Workbook(), Workbook().active
        ws.title = "DSE Results"
        ws.append(["Timestamp", "Stack_ID", "T", "D_in", "D_ffw", "D_attn", "Num_PE", "Freq_MHz", "BW_GBps", 
                   "FFW_Strat", "FFW_Tile", "Nonlin_Strat", "Nonlin_Tile", "Self_Strat", "Self_Tile", 
                   "Peak_SRAM_KB", "Total_MACs", "Global_Loads", "Global_Stores", "Local_Loads", "Local_Stores", "Intensity_MAC_Byte", "Latency_ms"])
        for cell in ws[1]: 
            cell.fill, cell.font, cell.alignment = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True), Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    for row in excel_rows_export: ws.append(row)
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 16
    wb.save(excel_path)
    print("✅ Đã lưu JSON và Excel thành công 100%!")

if __name__ == "__main__":
    run_multi_stack_dse()